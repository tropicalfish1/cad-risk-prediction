"""CAD prospective inference: extract features from test zips, predict, generate PPT."""
from __future__ import annotations

import csv, json, os, sys, time, zipfile
from collections import defaultdict
from pathlib import Path

import numpy as np
import pandas as pd
from scipy.signal import butter, filtfilt, find_peaks, welch
from xgboost import XGBClassifier
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor

sys.stdout.reconfigure(encoding='utf-8')

FEATURES = ["resting_hr_bpm", "sdnn_ms", "rmssd_ms", "lf_hf_ratio", "perfusion_index",
            "augmentation_index", "primary_reflected_peak_delay_ms", "st_deviation_mv",
            "qtc_ms", "pulse_transit_time_ms"]

ATTACH_DIR = Path(r"E:\data\dataset\attachments\attachments")
MAPPING_FILE = Path(r"E:\data\dataset\attachments\标签对应表.txt")
MODEL_FILE = Path(r"E:\data\cad_model_output\cad_xgboost.json")
OUTPUT_DIR = Path(r"E:\data\cad_inference_output")

def log(msg):
    print(f"[{time.strftime('%H:%M:%S')}] {msg}", flush=True)

def bandpass(x, fs, lo, hi):
    if len(x) < 30:
        return x
    b, a = butter(3, [lo / (fs / 2), min(hi / (fs / 2), 0.99)], btype="band")
    return filtfilt(b, a, x)

def ppg_features(ppg, fs):
    x = bandpass(np.asarray(ppg, dtype=float), fs, 0.5, 8)
    prom = max(np.std(x) * 0.35, 1e-9)
    peaks, _ = find_peaks(x, distance=int(fs * 0.35), prominence=prom)
    if len(peaks) < 4:
        return {k: np.nan for k in FEATURES[:7]}
    rr = np.diff(peaks) / fs
    rr = rr[(rr > 0.35) & (rr < 2.0)]
    hr = 60 / np.median(rr) if len(rr) else np.nan
    sdnn = np.std(rr, ddof=1) * 1000 if len(rr) > 1 else np.nan
    rmssd = np.sqrt(np.mean(np.diff(rr) ** 2)) * 1000 if len(rr) > 2 else np.nan
    f, psd = welch(1 / rr if len(rr) > 3 else np.array([0]), fs=4,
                   nperseg=min(64, max(4, len(rr))))
    def power(lo_, hi_):
        mask = (f >= lo_) & (f < hi_)
        return np.trapz(psd[mask], f[mask]) if mask.any() else 0.0
    hf, lf = power(0.15, 0.4), power(0.04, 0.15)
    pi = (np.percentile(x, 95) - np.percentile(x, 5)) / (abs(np.mean(x)) + 1e-9) * 100
    delays, ais = [], []
    for a, b in zip(peaks[:-1], peaks[1:]):
        trough = a + np.argmin(x[a:b])
        sub, _ = find_peaks(x[a + int(0.08 * fs):trough], prominence=prom * 0.1)
        if len(sub):
            r = a + int(0.08 * fs) + sub[-1]
            delays.append((r - a) / fs * 1000)
            ais.append((x[r] - x[a]) / (x[a] - x[trough] + 1e-9) * 100)
    return dict(resting_hr_bpm=hr, sdnn_ms=sdnn, rmssd_ms=rmssd,
                lf_hf_ratio=lf / (hf + 1e-9), perfusion_index=pi,
                augmentation_index=np.nanmedian(ais) if ais else np.nan,
                primary_reflected_peak_delay_ms=np.nanmedian(delays) if delays else np.nan)

def read_zip_sensor(zip_path):
    try:
        with zipfile.ZipFile(zip_path) as z:
            names = [x for x in z.namelist() if x != "info.json"]
            if not names:
                return None, None
            raw = z.read(names[0]).decode("utf-8", errors="ignore").splitlines()
    except Exception:
        return None, None

    header = raw[0].split('\t')
    ppg_time_idx = header.index('PPG_TIME') if 'PPG_TIME' in header else -1
    ppg9_idx = header.index('PPG9') if 'PPG9' in header else -1
    if ppg_time_idx < 0 or ppg9_idx < 0:
        return None, None

    ts_val_pairs = []
    for line in raw[1:]:
        parts = line.split('\t')
        if ppg_time_idx >= len(parts) or ppg9_idx >= len(parts):
            continue
        try:
            t = int(parts[ppg_time_idx])
            v = float(parts[ppg9_idx])
            if t > 0 and v != 0:
                ts_val_pairs.append((t, v))
        except (ValueError, IndexError):
            continue

    if len(ts_val_pairs) < 300:
        return None, None

    vals = np.array([p[1] for p in ts_val_pairs], dtype=float)
    ts_arr = np.array([p[0] for p in ts_val_pairs], dtype=float)

    unique_ts = np.unique(ts_arr)
    if len(unique_ts) < 2:
        return None, None
    median_block_gap = np.median(np.diff(unique_ts))
    samples_per_block = len(ts_val_pairs) / len(unique_ts)
    fs = samples_per_block * 1000.0 / median_block_gap if median_block_gap > 0 else 100.0
    fs = max(min(fs, 500), 10)

    return vals, fs

def load_mapping():
    mapping = {}
    with open(MAPPING_FILE, 'r', encoding='utf-8-sig') as f:
        reader = csv.reader(f)
        next(reader)
        for row in reader:
            data_id = row[0]
            external_id = row[3] if len(row) > 3 else ""
            parts = data_id.split('_')
            if len(parts) >= 2:
                mapping[parts[1]] = external_id
    return mapping

def main():
    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)

    log("Loading mapping table...")
    ts_to_ext = load_mapping()
    log(f"Loaded {len(ts_to_ext)} timestamp mappings")

    log("Loading XGBoost model...")
    model = XGBClassifier()
    model.load_model(str(MODEL_FILE))

    zip_files = sorted(f for f in os.listdir(ATTACH_DIR) if f.endswith('.zip'))
    log(f"Found {len(zip_files)} test zip files")

    results = []
    processed = 0
    errors = 0

    for zip_name in zip_files:
        zip_path = ATTACH_DIR / zip_name
        try:
            with zipfile.ZipFile(zip_path) as z:
                names = [x for x in z.namelist() if x != "info.json"]
                if not names:
                    errors += 1
                    continue
                fname = names[0].replace('.txt', '')
                parts = fname.split('_')
                if len(parts) < 3:
                    errors += 1
                    continue
                ts_str = parts[2]

            external_id = ts_to_ext.get(ts_str, "UNKNOWN")

            sig, fs = read_zip_sensor(zip_path)
            if sig is None or fs is None:
                errors += 1
                continue

            feats = ppg_features(sig, fs)
            feats["st_deviation_mv"] = np.nan
            feats["qtc_ms"] = np.nan
            feats["pulse_transit_time_ms"] = np.nan

            row = {f: feats.get(f, np.nan) for f in FEATURES}
            row["external_id"] = external_id
            row["zip_file"] = zip_name
            results.append(row)
            processed += 1

            if processed % 50 == 0:
                log(f"Processed {processed}/{len(zip_files)} zips ({errors} errors)")

        except Exception:
            errors += 1
            continue

    log(f"Feature extraction complete: {processed} successful, {errors} errors")

    if not results:
        log("No results to output!")
        return

    df = pd.DataFrame(results)

    df.to_csv(OUTPUT_DIR / "features_extracted.csv", index=False, encoding="utf-8-sig")
    log(f"Saved extracted features to {OUTPUT_DIR / 'features_extracted.csv'}")

    log("Running XGBoost inference...")
    from sklearn.impute import SimpleImputer
    imputer = SimpleImputer(strategy='median', add_indicator=True)
    imputer.fit(pd.read_csv(r"E:\data\cad_model_output\training_features.csv")[FEATURES])
    X = imputer.transform(df[FEATURES])
    probs = model.predict_proba(X)[:, 1]
    df["cad_probability"] = probs
    df["cad_prediction"] = (probs >= 0.5).astype(int)
    df["risk_level"] = pd.cut(probs, bins=[0, 0.3, 0.5, 0.7, 1.0],
                              labels=["低风险", "中风险", "较高风险", "高风险"])

    patient_agg = df.groupby("external_id").agg(
        avg_probability=("cad_probability", "mean"),
        max_probability=("cad_probability", "max"),
        record_count=("cad_probability", "count"),
        avg_hr=("resting_hr_bpm", "mean"),
        avg_sdnn=("sdnn_ms", "mean"),
        avg_rmssd=("rmssd_ms", "mean"),
        avg_lf_hf=("lf_hf_ratio", "mean"),
        avg_pi=("perfusion_index", "mean"),
        avg_ai=("augmentation_index", "mean"),
        avg_rpd=("primary_reflected_peak_delay_ms", "mean"),
    ).reset_index()
    patient_agg["cad_prediction"] = (patient_agg["avg_probability"] >= 0.5).astype(int)
    patient_agg["risk_level"] = pd.cut(patient_agg["avg_probability"],
                                        bins=[0, 0.3, 0.5, 0.7, 1.0],
                                        labels=["低风险", "中风险", "较高风险", "高风险"])

    df.to_csv(OUTPUT_DIR / "per_record_predictions.csv", index=False, encoding="utf-8-sig")
    patient_agg.to_csv(OUTPUT_DIR / "patient_level_predictions.csv", index=False, encoding="utf-8-sig")
    log(f"Saved per-record predictions: {OUTPUT_DIR / 'per_record_predictions.csv'}")
    log(f"Saved patient-level predictions: {OUTPUT_DIR / 'patient_level_predictions.csv'}")

    log(f"\n{'='*60}")
    log(f"Inference Summary:")
    log(f"  Total records: {len(df)}")
    log(f"  Unique patients: {len(patient_agg)}")
    log(f"  Predicted CAD positive: {(patient_agg['cad_prediction']==1).sum()}")
    log(f"  Predicted CAD negative: {(patient_agg['cad_prediction']==0).sum()}")
    log(f"  Risk distribution:")
    for level in ["低风险", "中风险", "较高风险", "高风险"]:
        count = int((patient_agg["risk_level"] == level).sum())
        log(f"    {level}: {count}")
    log(f"{'='*60}")

    generate_ppt(patient_agg, df)

def _add_text_box(slide, left, top, width, height, text, size=12, bold=False, color=None):
    txBox = slide.shapes.add_textbox(left, top, width, height)
    tf = txBox.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = text
    p.font.size = Pt(size)
    if bold:
        p.font.bold = True
    if color:
        p.font.color.rgb = color

def _set_cell(cell, text, bold=False, bg=None, font_color=None):
    cell.text = str(text)
    for p in cell.text_frame.paragraphs:
        p.font.size = Pt(9)
        if bold:
            p.font.bold = True
        if font_color:
            p.font.color.rgb = font_color
    if bg:
        from pptx.oxml.ns import qn
        tcPr = cell._tc.get_or_add_tcPr()
        solidFill = tcPr.makeelement(qn('a:solidFill'), {})
        srgbClr = solidFill.makeelement(qn('a:srgbClr'), {'val': '%02x%02x%02x' % (bg[0] if isinstance(bg, tuple) else bg.red, bg[1] if isinstance(bg, tuple) else bg.green, bg[2] if isinstance(bg, tuple) else bg.blue)})
        solidFill.append(srgbClr)
        tcPr.append(solidFill)

def generate_ppt(patient_agg, per_record_df):
    log("Generating PowerPoint presentation...")
    prs = Presentation()
    prs.slide_width = Inches(13.33)
    prs.slide_height = Inches(7.5)

    slide = prs.slides.add_slide(prs.slide_layouts[6])
    _add_text_box(slide, Inches(1), Inches(1.5), Inches(11), Inches(1.5),
                  "CAD AI辅助筛查模型", 36, bold=True, color=RGBColor(0, 51, 102))
    _add_text_box(slide, Inches(1), Inches(3.2), Inches(11), Inches(1),
                  "前瞻性验证数据集预测结果报告", 24, color=RGBColor(80, 80, 80))
    _add_text_box(slide, Inches(1), Inches(4.5), Inches(11), Inches(1),
                  f"基于PPG信号特征的XGBoost模型 | 样本量: {len(patient_agg)}例患者 / {len(per_record_df)}条记录",
                  16, color=RGBColor(120, 120, 120))

    slide2 = prs.slides.add_slide(prs.slide_layouts[6])
    _add_text_box(slide2, Inches(0.5), Inches(0.3), Inches(12), Inches(0.8),
                  "一、模型概况", 28, bold=True, color=RGBColor(0, 51, 102))
    info_lines = [
        "模型类型: XGBoost分类器 (max_depth=4, n_estimators=300)",
        "训练集: 106,973例样本 (阳性6,340例)",
        "交叉验证AUC: 0.6824",
        "输入特征(7个PPG特征+3个ECG空值):",
        "  resting_hr_bpm, sdnn_ms, rmssd_ms, lf_hf_ratio,",
        "  perfusion_index, augmentation_index, primary_reflected_peak_delay_ms,",
        "  st_deviation_mv(NaN), qtc_ms(NaN), pulse_transit_time_ms(NaN)",
        "备注: 前瞻性数据中无ECG信号，后三个特征为NaN",
    ]
    for i, line in enumerate(info_lines):
        _add_text_box(slide2, Inches(0.8), Inches(1.3 + i * 0.5), Inches(11), Inches(0.5),
                      line, 14, color=RGBColor(40, 40, 40))

    slide3 = prs.slides.add_slide(prs.slide_layouts[6])
    _add_text_box(slide3, Inches(0.5), Inches(0.3), Inches(12), Inches(0.8),
                  "二、预测结果总览", 28, bold=True, color=RGBColor(0, 51, 102))

    total_patients = len(patient_agg)
    pos_count = int((patient_agg["cad_prediction"] == 1).sum())
    neg_count = int((patient_agg["cad_prediction"] == 0).sum())
    mean_prob = patient_agg["avg_probability"].mean()

    risk_counts = patient_agg["risk_level"].value_counts()
    stats_lines = [
        f"总患者数: {total_patients}",
        f"预测阳性(CAD风险): {pos_count} ({pos_count / total_patients * 100:.1f}%)",
        f"预测阴性(低风险): {neg_count} ({neg_count / total_patients * 100:.1f}%)",
        f"平均预测概率: {mean_prob:.4f}",
        "",
        "风险等级分布:",
    ]
    for level in ["低风险", "中风险", "较高风险", "高风险"]:
        c = int(risk_counts.get(level, 0))
        stats_lines.append(f"  {level}: {c}例 ({c / total_patients * 100:.1f}%)")

    for i, line in enumerate(stats_lines):
        _add_text_box(slide3, Inches(0.8), Inches(1.3 + i * 0.45), Inches(11), Inches(0.45),
                      line, 14, color=RGBColor(40, 40, 40))

    slide4 = prs.slides.add_slide(prs.slide_layouts[6])
    _add_text_box(slide4, Inches(0.5), Inches(0.3), Inches(12), Inches(0.8),
                  "三、关键特征统计", 28, bold=True, color=RGBColor(0, 51, 102))

    ppg_feats = ["resting_hr_bpm", "sdnn_ms", "rmssd_ms", "lf_hf_ratio",
                 "perfusion_index", "augmentation_index", "primary_reflected_peak_delay_ms"]
    feat_stats = per_record_df[ppg_feats].describe().T
    feat_stats["median"] = per_record_df[ppg_feats].median()
    table_data = []
    for feat in ppg_feats:
        if feat in feat_stats.index:
            row = feat_stats.loc[feat]
            table_data.append([feat, f"{row.get('mean', 0):.2f}", f"{row.get('median', 0):.2f}",
                               f"{row.get('std', 0):.2f}", f"{row.get('min', 0):.2f}",
                               f"{row.get('max', 0):.2f}"])

    tbl = slide4.shapes.add_table(len(table_data) + 1, 6,
                                   Inches(0.5), Inches(1.3), Inches(12), Inches(5)).table
    headers = ["特征", "均值", "中位数", "标准差", "最小值", "最大值"]
    for j, h in enumerate(headers):
        _set_cell(tbl.cell(0, j), h, bold=True, bg=RGBColor(0, 51, 102), font_color=RGBColor(255, 255, 255))
    for i, row_data in enumerate(table_data):
        bg = RGBColor(240, 245, 250) if i % 2 == 0 else None
        for j, val in enumerate(row_data):
            _set_cell(tbl.cell(i + 1, j), val, bg=bg)

    slide5 = prs.slides.add_slide(prs.slide_layouts[6])
    _add_text_box(slide5, Inches(0.5), Inches(0.3), Inches(12), Inches(0.8),
                  "四、患者级预测结果(按概率降序，前30名)", 28, bold=True, color=RGBColor(0, 51, 102))

    top = patient_agg.sort_values("avg_probability", ascending=False).head(30)
    tbl2 = slide5.shapes.add_table(len(top) + 1, 7,
                                    Inches(0.3), Inches(1.2), Inches(12.7), Inches(5.5)).table

    headers2 = ["患者ID", "平均概率", "最高概率", "记录数", "平均HR", "预测结果", "风险等级"]
    for j, h in enumerate(headers2):
        _set_cell(tbl2.cell(0, j), h, bold=True, bg=RGBColor(0, 51, 102), font_color=RGBColor(255, 255, 255))
    for i, (_, r) in enumerate(top.iterrows()):
        pred_text = "阳性" if r["cad_prediction"] == 1 else "阴性"
        risk_text = str(r["risk_level"]) if pd.notna(r["risk_level"]) else "N/A"
        hr_val = f"{r['avg_hr']:.1f}" if pd.notna(r["avg_hr"]) else "N/A"
        row_data = [
            str(r["external_id"]),
            f"{r['avg_probability']:.4f}",
            f"{r['max_probability']:.4f}",
            str(int(r["record_count"])),
            hr_val,
            pred_text,
            risk_text,
        ]
        bg = RGBColor(255, 230, 230) if r["cad_prediction"] == 1 else RGBColor(230, 255, 230)
        for j, val in enumerate(row_data):
            _set_cell(tbl2.cell(i + 1, j), val, bg=bg)

    slide6 = prs.slides.add_slide(prs.slide_layouts[6])
    _add_text_box(slide6, Inches(0.5), Inches(0.3), Inches(12), Inches(0.8),
                  "五、结论与说明", 28, bold=True, color=RGBColor(0, 51, 102))
    min_p = patient_agg['avg_probability'].min()
    max_p = patient_agg['avg_probability'].max()
    conclusion_lines = [
        f"1. 共对{total_patients}例前瞻性患者进行CAD风险预测，其中{pos_count}例被判定为阳性({pos_count / total_patients * 100:.1f}%)",
        f"2. 模型使用PPG信号提取的7个生理特征(HRV、灌注指数、增强指数等)+3个ECG特征(填充值)",
        f"3. 平均预测概率为{mean_prob:.4f}，概率分布集中在{min_p:.3f}~{max_p:.3f}",
        "",
        "注意事项:",
        "  - 本模型基于回顾性数据训练(AUC=0.6824)，前瞻性验证结果仅供参考",
        "  - 前瞻性数据中无ECG信号，3个ECG相关特征使用训练集填充",
        "  - 预测阈值默认为0.5，可根据临床需求调整",
        "  - 建议结合临床检查结果综合判断",
    ]
    for i, line in enumerate(conclusion_lines):
        _add_text_box(slide6, Inches(0.8), Inches(1.3 + i * 0.5), Inches(11), Inches(0.5),
                      line, 14, color=RGBColor(40, 40, 40))

    ppt_path = OUTPUT_DIR / "CAD_前瞻性预测报告.pptx"
    prs.save(str(ppt_path))
    log(f"Saved PPT: {ppt_path}")

if __name__ == "__main__":
    main()
